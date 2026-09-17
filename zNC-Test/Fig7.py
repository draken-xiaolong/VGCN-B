#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Fig7：缩放攻击鲁棒性测试脚本（6个缩放比例）

比例参考图示：10%、50%、90%、130%、170%、210%
对应 scale 因子：0.1, 0.5, 0.9, 1.3, 1.7, 2.1

步骤与 Fig1 对齐，输出目录隔离在 `scale` 子目录。
"""

from pathlib import Path
import sys
from typing import List, Dict
import pickle
import shutil

import numpy as np

try:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass
# 强制刷新输出（确保实时写入日志）
import builtins
_original_print = builtins.print

def flush_print(*args, **kwargs):
    """带自动刷新的print函数"""
    _original_print(*args, **kwargs)
    sys.stdout.flush()

# 将print替换为flush_print以确保实时输出
builtins.print = flush_print


# 导入共享模块 
try:
    from fig_common import (
        PROJECT_ROOT, MODEL_PATH, CAT32_PATH, K_FOR_KNN,
        extract_features_20d, gdf_to_graph, load_improved_gat_model,
        load_cat32, features_to_matrix, calc_nc, extract_features_from_graph,
        convert_to_geojson, convert_geojsons_to_graphs
    ,
        generate_scramble_key, scramble_image, descramble_image,
        save_scramble_info, load_scramble_info
    )
except ImportError as e:
    print(f"无法导入 fig_common 模块: {e}")
    print("请确保 fig_common.py 在同一目录下")
    sys.exit(1)

try:
    import geopandas as gpd  # type: ignore
except Exception as exc:
    print("需要安装 geopandas: pip install geopandas fiona pyproj shapely")
    print("geopandas_import_error", exc)
    gpd = None  # type: ignore

try:
    import torch  # type: ignore
    from torch_geometric.data import Data  # type: ignore
except Exception as exc:
    print("需要安装 torch 和 torch-geometric")
    print("torch_import_error", exc)
    Data = None  # type: ignore

try:
    from sklearn.preprocessing import StandardScaler  # type: ignore
    from sklearn.neighbors import NearestNeighbors, kneighbors_graph  # type: ignore
except Exception as exc:
    print("需要安装 scikit-learn: pip install scikit-learn")
    print("sklearn_import_error", exc)
    StandardScaler = None  # type: ignore
    NearestNeighbors = None  # type: ignore
    kneighbors_graph = None  # type: ignore

try:
    import pandas as pd  # type: ignore
    # ⚠️ 关键：必须在任何matplotlib操作之前设置后端
    import matplotlib
    matplotlib.use('Agg')  # 使用非交互式后端
    
    # 配置中文字体（在设置后端之后）
    matplotlib.rcParams["font.sans-serif"] = [
        "SimHei", "Microsoft YaHei", "Microsoft JhengHei",
        "WenQuanYi Zen Hei", "Noto Sans CJK SC", "DejaVu Sans",
    ]
    matplotlib.rcParams["axes.unicode_minus"] = False
    
    import matplotlib.pyplot as plt  # type: ignore
    from PIL import Image  # type: ignore
except Exception as exc:
    print("需要安装 pandas, matplotlib, pillow")
    print("pandas_import_error", exc)
    pd = None  # type: ignore
    plt = None  # type: ignore
    Image = None  # type: ignore

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCRIPT_DIR = Path(__file__).resolve().parent

DIR_VECTOR = SCRIPT_DIR / 'vector-data'
DIR_VECTOR_GEOJSON = SCRIPT_DIR / 'vector-data-geojson'
DIR_VECTOR_GEOJSON_ATTACKED = SCRIPT_DIR / 'vector-data-geojson-attacked' / 'scale'
DIR_GRAPH = SCRIPT_DIR / 'vector-data-geojson-attacked-graph'
DIR_GRAPH_ORIGINAL = DIR_GRAPH / 'Original'
DIR_GRAPH_ATTACKED = DIR_GRAPH / 'Attacked' / 'scale'
DIR_ZEROWM = SCRIPT_DIR / 'vector-data-zerowatermark'
DIR_RESULTS = SCRIPT_DIR / 'NC-Results' / 'Fig7'

# MODEL_PATH和CAT32_PATH已从fig_common导入，无需重复定义
SCALE_FACTORS = [0.1, 0.5, 0.9, 1.3, 1.7, 2.1]


def step1_discover_inputs() -> List[Path]:
    print('[Step1] 扫描输入数据: ', DIR_VECTOR)
    if not DIR_VECTOR.exists():
        print('未找到目录: ', DIR_VECTOR)
        return []
    files: List[Path] = []
    files.extend(sorted(DIR_VECTOR.glob('*.shp')))
    files.extend(sorted(DIR_VECTOR.glob('*.geojson')))
    selected = files[:8]
    print('发现文件: ', [p.name for p in selected])
    return selected


def step2_convert_to_geojson(inputs: List[Path]) -> List[Path]:
    """按 convertToGeoJson 逻辑转为 GeoJSON。"""
    print('[Step2] 转换为 GeoJSON ->', DIR_VECTOR_GEOJSON)
    return convert_to_geojson(inputs, DIR_VECTOR_GEOJSON)
    for src in inputs:
        try:
            base = src.stem
            out_path = DIR_VECTOR_GEOJSON / f'{base}.geojson'
            gdf = gpd.read_file(src)
            if getattr(gdf, 'crs', None) and str(gdf.crs) != 'EPSG:4326':
                gdf = gdf.to_crs('EPSG:4326')
            gdf.to_file(out_path, driver='GeoJSON', encoding='utf-8')
            print(f'输出: {out_path.name} ({len(gdf)} 要素)')
            outputs.append(out_path)
        except Exception as exc:
            print('convert_error', src.name, exc)
    return outputs


def step3_generate_scale_attacks(original_geojsons: List[Path]) -> Dict[str, Dict[float, Path]]:
    print('[Step3] 生成缩放攻击 ->', DIR_VECTOR_GEOJSON_ATTACKED)
    if gpd is None:
        print('缺少 geopandas，无法生成攻击。')
        return {}
    from shapely.affinity import scale as shp_scale  # type: ignore
    DIR_VECTOR_GEOJSON_ATTACKED.mkdir(parents=True, exist_ok=True)
    outputs: Dict[str, Dict[float, Path]] = {}
    for src in original_geojsons:
        base = src.stem
        subdir = DIR_VECTOR_GEOJSON_ATTACKED / base
        if subdir.exists():
            shutil.rmtree(subdir)
        subdir.mkdir(parents=True, exist_ok=True)
        try:
            gdf = gpd.read_file(src)
        except Exception as exc:
            print('read_geojson_error', src.name, exc)
            continue
        
        # ✅ 使用全局中心作为缩放原点，与Fig8旋转策略一致
        bounds = gdf.total_bounds  # [minx, miny, maxx, maxy]
        global_center_x = (bounds[0] + bounds[2]) / 2
        global_center_y = (bounds[1] + bounds[3]) / 2
        global_center = (global_center_x, global_center_y)
        
        outputs[base] = {}
        for s in SCALE_FACTORS:
            attacked = gdf.copy()
            attacked['geometry'] = attacked['geometry'].apply(lambda geom: shp_scale(geom, s, s, origin=global_center))
            out_path = subdir / f'scale_{int(round(s*100))}pct.geojson'
            try:
                attacked.to_file(out_path, driver='GeoJSON')
                print(f'输出: {base}/{out_path.name}')
                outputs[base][s] = out_path
            except Exception as exc:
                print('attack_error', base, s, exc)
                continue
    return outputs


def step4_convert_to_graph(original_geojsons: List[Path], attacked_map: Dict[str, Dict[str, Path]]):
    """使用 fig_common 的标准函数转为图结构（KNN无向图 k=8，13维特征，与训练集一致）。"""
    print('[Step4] 转换为图结构 ->', DIR_GRAPH)
    
    if gpd is None or Data is None:
        print('缺少依赖，无法构图。')
        return
    
    # 使用共享函数批量转换（13维特征 + KNN无向图，与训练集一致）
    convert_geojsons_to_graphs(
        original_geojsons=original_geojsons,
        attacked_geojson_map=attacked_map,
        output_dir_original=DIR_GRAPH_ORIGINAL,
        output_dir_attacked=DIR_GRAPH_ATTACKED
    )


def step5_generate_zero_watermark():
    """用 VGAT 模型对 Original 生成零水印。"""
    print('[Step5] 生成零水印 ->', DIR_ZEROWM)
    
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    model, device = load_improved_gat_model(device)
    copyright_img = load_cat32()
    
    DIR_ZEROWM.mkdir(parents=True, exist_ok=True)
    
    cnt = 0
    for pkl in sorted(DIR_GRAPH_ORIGINAL.glob('*_graph.pkl')):
        try:
            with open(pkl, 'rb') as f:
                graph = pickle.load(f)
            
            base = pkl.stem.replace('_graph', '')
            
            # ⭐检查置乱信息文件是否已存在
            scramble_info_path = DIR_ZEROWM / f'{base}_scramble_info.json'
            if scramble_info_path.exists():
                # 使用现有的置乱密钥（固定不变）
                _, scramble_key, timestamp = load_scramble_info(scramble_info_path)
                print(f'使用现有置乱密钥 - 数据集: {base}, 密钥: {scramble_key}')
            else:
                # 首次生成置乱密钥
                scramble_key, timestamp = generate_scramble_key(base)
                print(f'生成新置乱密钥 - 数据集: {base}, 密钥: {scramble_key}')
            
            # ⭐置乱版权图像
            scrambled_copyright = scramble_image(copyright_img, scramble_key)
            
            # ⭐提取鲁棒特征并与置乱后的版权图像进行XOR生成零水印
            feat_mat = extract_features_from_graph(graph, model, device, copyright_img.shape)
            zwm = np.logical_xor(feat_mat, scrambled_copyright).astype(np.uint8)
            
            # ⭐保存零水印和置乱信息
            np.save(DIR_ZEROWM / f'{base}_watermark.npy', zwm)
            save_scramble_info(base, scramble_key, timestamp, DIR_ZEROWM / f'{base}_scramble_info.json')
            
            try:
                from PIL import Image
                Image.fromarray((zwm * 255).astype(np.uint8)).save(DIR_ZEROWM / f'{base}_watermark.png')
            except Exception:
                pass
            print('零水印已保存: ', f'{base}_watermark.npy')
            cnt += 1
        except Exception as exc:
            print('zero_watermark_error', pkl.name, exc)
            continue
    
    if cnt == 0:
        print('未找到 Original 图，无法生成零水印。')
    else:
        print(f'共生成零水印 {cnt} 个。')


def step6_evaluate_nc():
    print('[Step6] 评估NC ->', DIR_RESULTS)
    if not DIR_GRAPH_ATTACKED.exists():
        print('缺少攻击图目录: ', DIR_GRAPH_ATTACKED); print('请先运行 Step4 转换为图结构。'); return
    if Data is None or not MODEL_PATH.exists():
        if not MODEL_PATH.exists(): print('模型文件不存在: ', MODEL_PATH)
        else: print('缺少 torch/torch-geometric 依赖。')
        return
    # 检查依赖（pandas和plt已在文件开头导入）
    if pd is None or plt is None:
        print('缺少pandas/matplotlib依赖'); return
    # 使用共享函数加载模型
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    try:
        model, device = load_improved_gat_model(device)
    except Exception as exc:
        print(f'模型加载失败: {exc}')
        return

    DIR_RESULTS.mkdir(parents=True, exist_ok=True)
    copyright_img = load_cat32()

    rows = []
    for base_dir_path in sorted(DIR_GRAPH_ATTACKED.iterdir()):
        if not base_dir_path.is_dir(): continue
        base = base_dir_path.name
        zwm_path = DIR_ZEROWM / f'{base}_watermark.npy'
        scramble_info_path = DIR_ZEROWM / f'{base}_scramble_info.json'
        
        if not zwm_path.exists():
            print('缺少零水印，跳过: ', base)
            continue
        
        # ⭐加载置乱信息
        if not scramble_info_path.exists():
            print(f'警告: 缺少置乱信息文件，跳过: {base}')
            continue
        
        _, scramble_key, _ = load_scramble_info(scramble_info_path)
        zwm = np.load(zwm_path)
        
        for pkl in sorted(base_dir_path.glob('*_graph.pkl')):
            try:
                with open(pkl,'rb') as f: graph: Data = pickle.load(f)
                feat_mat = extract_features_from_graph(graph, model, device, copyright_img.shape)
                # ⭐从零水印提取置乱后的版权图像
                scrambled_extracted = np.logical_xor(zwm, feat_mat).astype(np.uint8)
                
                # ⭐使用置乱密钥恢复版权图像
                extracted = descramble_image(scrambled_extracted, scramble_key)
                
                # ⭐计算NC值
                nc = calc_nc(copyright_img, extracted)
                name = pkl.stem.replace('_graph','')
                pct = None
                for s in SCALE_FACTORS:
                    if f'scale_{int(round(s*100))}pct' in name: pct = s; break
                rows.append({'base': base, 'scale': pct if pct is not None else -1.0, 'nc': nc})
            except Exception as exc:
                print('nc_eval_error', pkl.name, exc)

    if not rows:
        print('没有评估结果。'); return

    df = pd.DataFrame(rows)
    hierarchical_rows = []
    base_names = sorted(df['base'].unique())
    for s in SCALE_FACTORS:
        label = f'{int(round(s*100))}%'
        hierarchical_rows.append({'缩放比例': label, 'VGAT': '', '类型': 'header'})
        for base in base_names:
            sub = df[(df['base']==base) & (df['scale']==s)]
            nc_value = sub['nc'].iloc[0] if len(sub)>0 else 0
            hierarchical_rows.append({'缩放比例': f'  {base}', 'VGAT': f'{nc_value:.6f}', '类型': 'data'})
        avg_nc = df[df['scale']==s]['nc'].mean()
        hierarchical_rows.append({'缩放比例': '  Average', 'VGAT': f'{avg_nc:.6f}', '类型': 'average'})

    hierarchical_df = pd.DataFrame(hierarchical_rows)
    csv_path = DIR_RESULTS / 'fig7_scale_nc.csv'
    hierarchical_df.to_csv(csv_path, index=False, encoding='utf-8-sig')
    xlsx_path = DIR_RESULTS / 'fig7_scale_nc.xlsx'
    try:
        with pd.ExcelWriter(xlsx_path, engine='openpyxl') as writer:
            hierarchical_df.to_excel(writer, sheet_name='NC结果', index=False)
            ws = writer.sheets['NC结果']
            ws.column_dimensions['A'].width = 20
            ws.column_dimensions['B'].width = 15
            from openpyxl.styles import Font, Alignment
            for idx, row in hierarchical_df.iterrows():
                if row['类型'] == 'header':
                    ws[f'A{idx+2}'].font = Font(bold=True)
                elif row['类型'] == 'average':
                    ws[f'A{idx+2}'].font = Font(bold=True, italic=True)
                    ws[f'B{idx+2}'].font = Font(bold=True, italic=True)
                elif row['类型'] == 'data':
                    ws[f'A{idx+2}'].alignment = Alignment(horizontal='left', indent=1)
    except Exception as e:
        print('Excel保存失败: ', e)
        hierarchical_df.to_excel(xlsx_path, index=False, engine='openpyxl')
    print('结果表保存: ', csv_path.name)
    print('Excel文件保存: ', xlsx_path.name)

    # 绘图
    plt.figure(figsize=(10,6))
    x = np.arange(len(SCALE_FACTORS))
    labels = [f'{int(round(s*100))}%' for s in SCALE_FACTORS]
    for base, sub in df.groupby('base'):
        y = []
        for s in SCALE_FACTORS:
            val = sub[sub['scale']==s]['nc']
            y.append(val.iloc[0] if len(val)>0 else np.nan)
        plt.plot(x, y, '-o', alpha=0.7, label=base)
    avg_vals = [df[df['scale']==s]['nc'].mean() for s in SCALE_FACTORS]
    plt.plot(x, avg_vals, 'k-o', linewidth=2.5, label='平均')
    plt.grid(True, alpha=0.3)
    plt.xticks(x, labels)
    plt.xlabel('缩放比例')
    plt.ylabel('NC')
    plt.title('Fig7：缩放攻击 的NC鲁棒性')
    plt.legend(loc='best', fontsize=8, ncol=2)
    fig_path = DIR_RESULTS / 'fig7_scale_nc.png'
    plt.tight_layout(); plt.savefig(fig_path, dpi=300, bbox_inches='tight'); plt.close()
    print('曲线图保存: ', fig_path.name)


def check_existing_files():
    print('[检查] 检查现有文件...')
    geojson_files = [p for p in DIR_VECTOR_GEOJSON.glob('*.geojson') if not p.name.startswith('._')]
    step2_done = len(geojson_files) > 0
    attacked_dirs = [d for d in DIR_VECTOR_GEOJSON_ATTACKED.glob('*') if not d.name.startswith('.')]
    step3_done = len(attacked_dirs) > 0 and all(len(list(ad.glob('*.geojson'))) > 0 for ad in attacked_dirs if ad.is_dir())
    original_graphs = [p for p in DIR_GRAPH_ORIGINAL.glob('*_graph.pkl') if not p.name.startswith('.')]
    attacked_graphs = [p for p in DIR_GRAPH_ATTACKED.glob('*/*_graph.pkl') if not p.name.startswith('.')]
    step4_done = len(original_graphs) > 0 and len(attacked_graphs) > 0
    watermark_files = [p for p in DIR_ZEROWM.glob('*_watermark.npy') if not p.name.startswith('.')]
    step5_done = len(watermark_files) > 0
    nc_files = [p for p in DIR_RESULTS.glob('*.csv') if not p.name.startswith('.')] + [p for p in DIR_RESULTS.glob('*.xlsx') if not p.name.startswith('.')] + [p for p in DIR_RESULTS.glob('*.png') if not p.name.startswith('.')]
    step6_done = len(nc_files) > 0
    print(f'  Step2 (GeoJSON): {"✓" if step2_done else "✗"} ({len(geojson_files)} files)')
    print(f'  Step3 (Attacked): {"✓" if step3_done else "✗"} ({len(attacked_dirs)} dirs)')
    print(f'  Step4 (Graphs): {"✓" if step4_done else "✗"} (Original: {len(original_graphs)}, Attacked: {len(attacked_graphs)})')
    print(f'  Step5 (Watermarks): {"✓" if step5_done else "✗"} ({len(watermark_files)} files)')
    print(f'  Step6 (NC Results): {"✓" if step6_done else "✗"} ({len(nc_files)} files)')
    return {'step2': step2_done, 'step3': step3_done, 'step4': step4_done, 'step5': step5_done, 'step6': step6_done}


def main():
    print('=== Fig7：缩放攻击鲁棒性测试 ===')
    existing = check_existing_files()
    if existing['step6']:
        print('\n[跳过] 检测到NC结果文件已存在，直接运行Step6...'); step6_evaluate_nc(); print('=== 完成 ==='); return
    if existing['step5'] and existing['step4']:
        print('\n[跳过] 检测到零水印与图结构已存在，直接运行Step6...'); step6_evaluate_nc(); print('=== 完成 ==='); return
    if existing['step4']:
        print('\n[跳过] 检测到图结构文件已存在，直接运行Step5-6...'); step5_generate_zero_watermark(); step6_evaluate_nc(); print('=== 完成 ==='); return
    if existing['step3']:
        print('\n[跳过] 检测到攻击文件已存在，直接运行Step4-6...')
        original_geojsons = [p for p in DIR_VECTOR_GEOJSON.glob('*.geojson') if not p.name.startswith('._')]
        attacked_map: Dict[str, Dict[float, Path]] = {}
        for attacked_dir in DIR_VECTOR_GEOJSON_ATTACKED.glob('*'):
            if attacked_dir.is_dir():
                base = attacked_dir.name; attacked_map[base] = {}
                for s in SCALE_FACTORS:
                    geojson_file = attacked_dir / f'scale_{int(round(s*100))}pct.geojson'
                    if geojson_file.exists(): attacked_map[base][s] = geojson_file
        step4_convert_to_graph(original_geojsons, attacked_map); step5_generate_zero_watermark(); step6_evaluate_nc(); print('=== 完成 ==='); return
    if existing['step2']:
        print('\n[跳过] 检测到GeoJSON文件已存在，直接运行Step3-6...')
        original_geojsons = [p for p in DIR_VECTOR_GEOJSON.glob('*.geojson') if not p.name.startswith('._')]
        attacked_map = step3_generate_scale_attacks(original_geojsons)
        step4_convert_to_graph(original_geojsons, attacked_map); step5_generate_zero_watermark(); step6_evaluate_nc(); print('=== 完成 ==='); return
    print('\n[从头开始] 运行所有步骤...')
    inputs = step1_discover_inputs(); original_geojsons = step2_convert_to_geojson(inputs)
    attacked_map = step3_generate_scale_attacks(original_geojsons)
    step4_convert_to_graph(original_geojsons, attacked_map); step5_generate_zero_watermark(); step6_evaluate_nc(); print('=== 完成 ===')
    sys.exit(0)


if __name__ == '__main__':
    main()
